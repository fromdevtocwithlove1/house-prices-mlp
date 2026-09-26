"""Check final DOCX/PDF against the source artifacts and report requirements."""

from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
SOURCE = REPORTS / "bao_cao_crisp_dm.md"
DOCX = REPORTS / "bao_cao_crisp_dm.docx"
PDF = REPORTS / "bao_cao_crisp_dm.pdf"
FIGURES = [
    "target_histograms.png",
    "missing_values.png",
    "numeric_scatter.png",
    "e0_rmse.png",
    "e0_actual_predicted.png",
    "e0_residuals.png",
]


def digest(data):
    return hashlib.sha256(data).hexdigest()


source = SOURCE.read_text(encoding="utf-8")
body = source[source.index("## 1. MỞ ĐẦU") : source.index("## 9. TÀI LIỆU THAM KHẢO")]
body = re.sub(r"!\[[^]]*\]\([^)]*\)", "", body)
body = re.sub(r"(?m)^\*Hình.*$|^\|.*$|^#.*$", "", body)
words = len(body.split())
assert 4000 <= words <= 5000, words

with ZipFile(DOCX) as archive:
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 6, media
    embedded_hashes = {digest(archive.read(name)) for name in media}
    original_hashes = {digest((ROOT / "outputs" / "figures" / name).read_bytes()) for name in FIGURES}
    assert embedded_hashes == original_hashes, "Embedded PNG bytes differ from source figures"

document = Document(DOCX)
captions = [p.text for p in document.paragraphs if p.style.name == "Caption" and p.text.startswith("Hình ")]
assert len(captions) == 6, captions
for number, caption in enumerate(captions, 1):
    assert caption.startswith(f"Hình {number}. ")
    assert "Nguồn: Kết quả thực nghiệm của dự án." in caption

tables = document.tables
assert len(tables) == 1
table = tables[0]
assert len(table.rows) == 6 and len(table.columns) == 7
with (ROOT / "outputs" / "metrics" / "experiments.csv").open(encoding="utf-8", newline="") as stream:
    experiments = list(csv.DictReader(stream))
assert len(experiments) == 5
for index, row in enumerate(experiments, 1):
    cells = [cell.text for cell in table.rows[index].cells]
    assert cells[0] == row["experiment_name"]
    assert cells[5] == row["best_epoch"]
    assert cells[6] == f'{float(row["best_val_rmse"]):.6f}'.replace(".", ",")

reader = PdfReader(PDF)
assert len(reader.pages) == 14, len(reader.pages)
texts = [page.extract_text() or "" for page in reader.pages]
all_text = "\n".join(texts)
for snippet in ("MỤC LỤC", "DỰ ĐOÁN GIÁ NHÀ", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC", "Kết quả thực nghiệm của dự án", "0,133316", "0.13543"):
    assert snippet in all_text, snippet
for number in range(1, 7):
    assert f"Hình {number}." in all_text, number
for exp in ("E0", "E1", "E2", "E3", "E4"):
    assert exp in texts[9], f"Missing table row {exp} on physical page 10"

headings = [
    "1. MỞ ĐẦU",
    "2. BUSINESS UNDERSTANDING — HIỂU BÀI TOÁN",
    "3. DATA UNDERSTANDING — HIỂU DỮ LIỆU",
    "4. DATA PREPARATION — CHUẨN BỊ DỮ LIỆU",
    "5. MODELING — XÂY DỰNG MÔ HÌNH",
    "6. EVALUATION — ĐÁNH GIÁ",
    "7. DEPLOYMENT — TRIỂN KHAI",
    "8. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN",
    "9. TÀI LIỆU THAM KHẢO",
    "10. PHỤ LỤC — LỆNH TÁI LẬP",
]
toc_lines = [line.strip() for line in texts[1].splitlines() if line.strip()]
for heading in headings:
    page_index = next(i for i, value in enumerate(texts[2:], 2) if heading in value)
    printed_page = page_index - 1
    matches = [line for line in toc_lines if heading in line]
    assert len(matches) == 1, (heading, matches)
    assert re.search(rf"\b{printed_page}$", matches[0]), (heading, matches[0], printed_page)

pdf_image_pages = [i + 1 for i, page in enumerate(reader.pages) if len(page.images) > 0]
assert len(pdf_image_pages) >= 5, pdf_image_pages
print(f"PASS words={words}, embedded_png={len(media)}, captions={len(captions)}, experiments=5")
print(f"PASS pdf_pages={len(reader.pages)}, pdf_image_pages={pdf_image_pages}, toc_headings={len(headings)}")
print(f"DOCX={DOCX.stat().st_size} bytes PDF={PDF.stat().st_size} bytes")
