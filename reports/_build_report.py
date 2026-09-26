"""Build the requested Vietnamese DOCX from the project-grounded report source."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
SOURCE = REPORTS / "bao_cao_crisp_dm.md"
OUTPUT = REPORTS / "bao_cao_crisp_dm.docx"


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        tag = "w:" + side
        edge = borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), color)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def suppress_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    p_pr.append(borders)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    elem = OxmlElement("w:tblHeader")
    elem.set(qn("w:val"), "true")
    tr_pr.append(elem)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    text_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    text_run._r.append(instr)
    separator_run = paragraph.add_run()
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separator_run._r.append(separator)
    if placeholder:
        paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))
    pg_num.set(qn("w:fmt"), "decimal")


def setup_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    for key, size, before, after in (
        ("Title", 19, 0, 14),
        ("Heading 1", 14, 18, 9),
        ("Heading 2", 12, 12, 6),
    ):
        style = styles[key]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10.5)
    caption.font.italic = True
    caption.font.bold = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(9)
    for name in ("TOC 1", "TOC 2"):
        if name in styles:
            style = styles[name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)


def plain_markup(value):
    return value.replace("`", "").replace("**", "")


def add_figure(doc, source_line):
    match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", source_line)
    if not match:
        raise ValueError(source_line)
    image_path = (REPORTS / match.group(2)).resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    with Image.open(image_path) as img:
        width_px, height_px = img.size
    max_width = 15.5
    max_height = 10.6 if image_path.name == "e0_actual_predicted.png" else 15.0
    width_cm = min(max_width, max_height * width_px / height_px)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(image_path), width=Cm(width_cm))


def add_table(doc, lines):
    rows = [[plain_markup(c.strip()) for c in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) < 3:
        raise ValueError("Expected header, delimiter, and data")
    data = [rows[0], *rows[2:]]
    if any(len(row) != len(data[0]) for row in data):
        raise ValueError("Inconsistent table columns")
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(v) for v in (2.0, 2.25, 2.2, 2.1, 2.1, 2.25, 3.1)]
    for ri, values in enumerate(data):
        row = table.rows[ri]
        prevent_row_split(row)
        if ri == 0:
            repeat_header(row)
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = widths[ci]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if ri == 0:
                set_cell_fill(cell, "DCE7F2")
            elif ri % 2 == 0:
                set_cell_fill(cell, "F5F8FB")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci in (0, 1, 2) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.1)
            run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    raw = SOURCE.read_text(encoding="utf-8")
    lines = raw.splitlines()
    doc = Document()
    setup_styles(doc)
    setup_section(doc.sections[0])

    # Cover page.
    doc.add_paragraph().paragraph_format.space_after = Pt(85)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("BÁO CÁO TIỂU LUẬN")
    suppress_paragraph_border(p)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("DỰ ĐOÁN GIÁ NHÀ BẰNG MLP PYTORCH")
    p.add_run().add_break()
    p.add_run("THEO CRISP-DM")
    suppress_paragraph_border(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(65)
    for value in ("Sinh viên: ............................................................",
                  "Lớp: ....................................................................",
                  "Giảng viên hướng dẫn: ..............................................."):
        p = doc.add_paragraph(value)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(3)
        p.paragraph_format.space_after = Pt(16)
    doc.add_paragraph().paragraph_format.space_after = Pt(115)
    p = doc.add_paragraph("Tháng 9 năm 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Separate TOC page. Word will populate the native field before PDF export.
    toc_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(toc_section)
    toc_section.header.is_linked_to_previous = False
    toc_section.footer.is_linked_to_previous = False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-1" \\h \\z \\u', "Đang cập nhật mục lục")

    # Restart visible page numbering at the introduction.
    body_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(body_section)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    set_page_number_start(body_section, 1)
    hp = body_section.header.paragraphs[0]
    hp.text = "Dự đoán giá nhà bằng MLP PyTorch theo CRISP-DM"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.name = "Times New Roman"
    hp.runs[0].font.size = Pt(9)
    fp = body_section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Trang ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    add_field(fp, "PAGE", "1")

    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. MỞ ĐẦU"))
    i = start
    headings = []
    figures = []
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            title = plain_markup(line[3:])
            doc.add_paragraph(title, style="Heading 1")
            headings.append(title)
            i += 1
            continue
        if line.startswith("!["):
            add_figure(doc, line)
            figures.append(line)
            i += 1
            continue
        if line.startswith("*Hình ") or line.startswith("*Bảng "):
            p = doc.add_paragraph(line.strip("*"), style="Caption")
            if line.startswith("*Bảng "):
                p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("`") and line.endswith("`") and line.count("`") == 2:
            p = doc.add_paragraph(plain_markup(line))
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            i += 1
            continue
        p = doc.add_paragraph(plain_markup(line))
        p.paragraph_format.first_line_indent = Cm(0.65)
        if line.startswith("[") and "] " in line[:5]:
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.left_indent = Cm(0.45)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(11)
        i += 1

    if len(figures) != 6:
        raise ValueError(f"Expected six required figures, found {len(figures)}")
    if len(headings) != 10:
        raise ValueError(f"Expected ten numbered headings, found {len(headings)}")
    doc.core_properties.title = "Báo cáo dự đoán giá nhà bằng MLP PyTorch theo CRISP-DM"
    doc.core_properties.subject = "Báo cáo dự án học máy theo sáu giai đoạn CRISP-DM"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"headings={len(headings)} figures={len(figures)}")


if __name__ == "__main__":
    main()
